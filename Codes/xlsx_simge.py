import io
import posixpath
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

KALIP = re.compile(r"(_|\^)\{([^}]*)\}")
ILISKI = re.compile(r"<Relationship\b[^>]*/>")


def temiz_paket(dosya):
    giris = zipfile.ZipFile(dosya)
    adlar = set(giris.namelist())
    cikis = io.BytesIO()
    with zipfile.ZipFile(cikis, "w", zipfile.ZIP_DEFLATED) as z:
        for ad in giris.namelist():
            veri = giris.read(ad)
            if ad.endswith(".rels"):
                taban = posixpath.dirname(posixpath.dirname(ad))
                metin = veri.decode("utf-8")

                def tut(m):
                    hedef = re.search(r'Target="([^"]*)"', m.group(0)).group(1)
                    if 'TargetMode="External"' in m.group(0) or hedef.startswith("/"):
                        return m.group(0)
                    return m.group(0) if posixpath.normpath(posixpath.join(taban, hedef)) in adlar else ""

                veri = ILISKI.sub(tut, metin).encode("utf-8")
            z.writestr(ad, veri)
    cikis.seek(0)
    return cikis


def yazi(font, dikey=None):
    return InlineFont(rFont=font.name, sz=font.sz, b=font.b, i=font.i, vertAlign=dikey)


def zengin(metin, font):
    parcalar = []
    son = 0
    for m in KALIP.finditer(metin):
        if m.start() > son:
            parcalar.append(TextBlock(yazi(font), metin[son:m.start()]))
        parcalar.append(TextBlock(yazi(font, "subscript" if m.group(1) == "_" else "superscript"), m.group(2)))
        son = m.end()
    if son < len(metin):
        parcalar.append(TextBlock(yazi(font), metin[son:]))
    return CellRichText(parcalar)


dosya = Path(sys.argv[1])
wb = load_workbook(temiz_paket(dosya), rich_text=True)
n = 0
for ws in wb.worksheets:
    for satir in ws.iter_rows():
        for h in satir:
            if isinstance(h.value, str) and KALIP.search(h.value):
                h.value = zengin(h.value, h.font)
                n += 1
wb.save(dosya)
print(dosya.name, n, "hucrede alt/ust simge")
